import io

from openpyxl import load_workbook


def add(client, cookies, amount, category, date, description="x", location=None):
    resp = client.post(
        "/api/expenses",
        headers=cookies,
        json={"amount": amount, "description": description, "category": category, "date": date, "location": location},
    )
    assert resp.status_code == 201, resp.text
    return resp


def test_export_requires_authentication(client):
    resp = client.get("/api/export/expenses.xlsx")
    assert resp.status_code == 401


def test_export_returns_xlsx_headers(client, alice):
    resp = client.get("/api/export/expenses.xlsx", headers=alice)
    assert resp.status_code == 200
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'attachment; filename=expenses.xlsx' in resp.headers["content-disposition"]


def test_export_contains_expense_rows(client, alice):
    add(client, alice, 42.5, "Food", "2026-08-01", description="Groceries", location="Trader Joe's")
    add(client, alice, 10, "Transport", "2026-08-02", description="Bus")

    resp = client.get("/api/export/expenses.xlsx", headers=alice)
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active

    assert [c.value for c in ws[1]] == ["Date", "Category", "Description", "Location", "Amount (ILS)"]
    assert [c.value for c in ws[2]] == ["2026-08-01", "Food", "Groceries", "Trader Joe's", 42.5]
    # openpyxl round-trips an empty string as None and a whole-number float
    # as int; both are indistinguishable from the intended values in Excel.
    assert [c.value for c in ws[3]] == ["2026-08-02", "Transport", "Bus", None, 10]
    assert ws.cell(row=4, column=1).value == "Total"
    assert ws.cell(row=4, column=5).value == "=SUM(E2:E3)"


def test_export_empty_has_no_total_row(client, alice):
    resp = client.get("/api/export/expenses.xlsx", headers=alice)
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws.max_row == 1


def test_export_is_scoped_per_user(client, alice, bob):
    add(client, alice, 100, "Food", "2026-08-01")
    add(client, bob, 200, "Food", "2026-08-01")

    resp = client.get("/api/export/expenses.xlsx", headers=alice)
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    amounts = [ws.cell(row=r, column=5).value for r in range(2, ws.max_row) if ws.cell(row=r, column=1).value != "Total"]
    assert amounts == [100.0]
