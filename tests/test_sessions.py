def add(client, cookies, **overrides):
    payload = {
        "amount": 12.5,
        "description": "Coffee",
        "category": "Food",
        "date": "2026-08-07",
    }
    payload.update(overrides)
    resp = client.post("/api/expenses", headers=cookies, json=payload)
    assert resp.status_code == 201, resp.text
    return resp.json()


def test_new_account_has_a_standard_session(client, alice):
    sessions = client.get("/api/sessions", headers=alice).json()
    assert len(sessions) == 1
    assert sessions[0]["name"] == "Standard"
    assert sessions[0]["currency"] == "ILS"
    assert sessions[0]["is_standard"] is True
    assert sessions[0]["is_active"] is True
    assert sessions[0]["status"] == "open"


def test_create_session_auto_activates_it(client, alice):
    resp = client.post("/api/sessions", headers=alice, json={"name": "Japan Trip", "currency": "jpy"})
    assert resp.status_code == 201
    body = resp.json()
    assert body["name"] == "Japan Trip"
    assert body["currency"] == "JPY"  # normalized to uppercase
    assert body["is_active"] is True
    assert body["is_standard"] is False

    sessions = client.get("/api/sessions", headers=alice).json()
    assert len(sessions) == 2
    active = [s for s in sessions if s["is_active"]]
    assert len(active) == 1
    assert active[0]["name"] == "Japan Trip"


def test_invalid_currency_rejected(client, alice):
    resp = client.post("/api/sessions", headers=alice, json={"name": "Trip", "currency": "12"})
    assert resp.status_code == 422


def test_expenses_are_isolated_between_sessions(client, alice):
    add(client, alice, description="ILS coffee")
    client.post("/api/sessions", headers=alice, json={"name": "Japan Trip", "currency": "JPY"})

    japan_expenses = client.get("/api/expenses?month=2026-08", headers=alice).json()
    assert japan_expenses == []

    add(client, alice, description="Ramen")
    japan_expenses = client.get("/api/expenses?month=2026-08", headers=alice).json()
    assert len(japan_expenses) == 1
    assert japan_expenses[0]["description"] == "Ramen"

    sessions = client.get("/api/sessions", headers=alice).json()
    standard_id = next(s["id"] for s in sessions if s["is_standard"])
    client.post(f"/api/sessions/{standard_id}/activate", headers=alice)

    standard_expenses = client.get("/api/expenses?month=2026-08", headers=alice).json()
    assert len(standard_expenses) == 1
    assert standard_expenses[0]["description"] == "ILS coffee"


def test_budgets_are_isolated_between_sessions(client, alice):
    client.put("/api/budgets/Food", headers=alice, json={"monthly_limit": 300})
    client.post("/api/sessions", headers=alice, json={"name": "Japan Trip", "currency": "JPY"})

    assert client.get("/api/budgets", headers=alice).json() == {}

    client.put("/api/budgets/Food", headers=alice, json={"monthly_limit": 50000})
    assert client.get("/api/budgets", headers=alice).json() == {"Food": 50000.0}

    sessions = client.get("/api/sessions", headers=alice).json()
    standard_id = next(s["id"] for s in sessions if s["is_standard"])
    client.post(f"/api/sessions/{standard_id}/activate", headers=alice)
    assert client.get("/api/budgets", headers=alice).json() == {"Food": 300.0}


def test_standard_session_cannot_be_closed(client, alice):
    sessions = client.get("/api/sessions", headers=alice).json()
    standard_id = sessions[0]["id"]
    resp = client.post(f"/api/sessions/{standard_id}/close", headers=alice)
    assert resp.status_code == 400


def test_closed_session_rejects_new_expenses(client, alice):
    created = client.post(
        "/api/sessions", headers=alice, json={"name": "Japan Trip", "currency": "JPY"}
    ).json()
    client.post(f"/api/sessions/{created['id']}/close", headers=alice)

    resp = client.post(
        "/api/expenses",
        headers=alice,
        json={"amount": 10, "description": "x", "category": "Food", "date": "2026-08-07"},
    )
    assert resp.status_code == 400


def test_reopen_allows_expenses_again(client, alice):
    created = client.post(
        "/api/sessions", headers=alice, json={"name": "Japan Trip", "currency": "JPY"}
    ).json()
    client.post(f"/api/sessions/{created['id']}/close", headers=alice)
    client.post(f"/api/sessions/{created['id']}/reopen", headers=alice)

    resp = client.post(
        "/api/expenses",
        headers=alice,
        json={"amount": 10, "description": "x", "category": "Food", "date": "2026-08-07"},
    )
    assert resp.status_code == 201


def test_sessions_are_private_per_user(client, alice, bob):
    client.post("/api/sessions", headers=alice, json={"name": "Alice's Trip", "currency": "USD"})
    bob_sessions = client.get("/api/sessions", headers=bob).json()
    assert len(bob_sessions) == 1
    assert bob_sessions[0]["name"] == "Standard"


def test_cannot_activate_another_users_session(client, alice, bob):
    alice_session = client.post(
        "/api/sessions", headers=alice, json={"name": "Alice's Trip", "currency": "USD"}
    ).json()
    resp = client.post(f"/api/sessions/{alice_session['id']}/activate", headers=bob)
    assert resp.status_code == 404


def test_sessions_require_authentication(client):
    resp = client.get("/api/sessions")
    assert resp.status_code == 401


def test_export_header_uses_session_currency(client, alice):
    import io

    from openpyxl import load_workbook

    client.post("/api/sessions", headers=alice, json={"name": "Japan Trip", "currency": "JPY"})
    add(client, alice, description="Ramen", amount=1500)

    resp = client.get("/api/export/expenses.xlsx", headers=alice)
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert [c.value for c in ws[1]] == ["Date", "Category", "Description", "Location", "Amount (JPY)"]
